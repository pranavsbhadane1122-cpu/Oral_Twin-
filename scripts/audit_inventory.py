"""Inventory of a raw dataset folder. READ-ONLY on data/raw.

Reports what is actually on disk: the folder tree with counts, images per
category, image size distribution, unreadable files, exact-duplicate groups by
MD5, and every metadata file with whether it parses.

Usage:
  python scripts/audit_inventory.py mio piyarathne_oral
"""

import hashlib
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from PIL import Image  # noqa: E402

from src.utils.config import load_config  # noqa: E402

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff", ".jfif"}
META_EXTS = {".csv", ".json", ".xlsx", ".xls", ".txt", ".md", ".yaml", ".yml"}


def md5_of(path, chunk=1 << 20):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(chunk), b""):
            h.update(block)
    return h.hexdigest()


def parse_check(path):
    """Does this metadata file parse? Returns a short description."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".json":
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                keys = list(data)[:8]
                sizes = {k: (len(data[k]) if hasattr(data[k], "__len__") else "scalar")
                         for k in keys}
                return f"OK, dict with top-level keys {sizes}"
            return f"OK, {type(data).__name__} of length {len(data)}"
        if suffix == ".csv":
            import csv
            with open(path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                header = next(reader, [])
                rows = sum(1 for _ in reader)
            return f"OK, {rows} rows, columns: {header}"
        if suffix in (".xlsx", ".xls"):
            import openpyxl
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for sheet in book.worksheets:
                header = next(sheet.iter_rows(max_row=1, values_only=True), ())
                header = [str(h) for h in header if h is not None][:10]
                parts.append(f"{sheet.title}: {sheet.max_row} rows, {header}")
            book.close()
            return "OK, sheets -> " + " | ".join(parts)
        return "not parsed (plain text)"
    except Exception as exc:
        return f"PARSE FAILED: {exc!r}"


def audit(root, name, log):
    log(f"\n{'=' * 78}\nDATASET: {name}   ({root})\n{'=' * 78}")
    if not root.exists():
        log("  MISSING")
        return

    images, metas, others = [], [], []
    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        suffix = p.suffix.lower()
        if suffix in IMAGE_EXTS:
            images.append(p)
        elif suffix in META_EXTS:
            metas.append(p)
        else:
            others.append(p)

    log(f"\nfiles: {len(images)} images, {len(metas)} metadata, {len(others)} other")

    log("\nfolder tree (folders holding images):")
    per_dir = Counter(p.parent for p in images)
    for directory in sorted(per_dir):
        log(f"  {per_dir[directory]:>6}  {directory.relative_to(root)}")

    log("\nimage extensions:")
    for ext, n in Counter(p.suffix.lower() for p in images).most_common():
        log(f"  {ext}: {n}")

    # reuse the cached hashes when audit_hashes.py has already read the 9 GB
    cache_path = (PROJECT_ROOT / load_config()["paths"]["processed"]
                  / "audit_cache" / "hashes.json")
    cached = {}
    if cache_path.exists():
        blob = json.loads(cache_path.read_text(encoding="utf-8"))
        if name in blob:
            cached = {e["path"]: e["md5"] for e in blob[name]["entries"]}

    sizes, corrupt, md5s = Counter(), [], defaultdict(list)
    for p in images:
        rel = str(p.relative_to(root)).replace("\\", "/")
        try:
            with Image.open(p) as im:
                sizes[im.size] += 1
        except Exception as exc:
            corrupt.append((p.relative_to(root), repr(exc)))
            continue
        digest = cached.get(rel) or md5_of(p)
        md5s[digest].append(p.relative_to(root))

    log(f"\nimage size distribution ({len(sizes)} distinct sizes), top 12:")
    for (w, h), n in sizes.most_common(12):
        log(f"  {w}x{h}: {n}")
    if sizes:
        widths = sorted(w for (w, h), n in sizes.items() for _ in range(n))
        heights = sorted(h for (w, h), n in sizes.items() for _ in range(n))
        mid = len(widths) // 2
        log(f"  width  min {widths[0]} median {widths[mid]} max {widths[-1]}")
        log(f"  height min {heights[0]} median {heights[mid]} max {heights[-1]}")

    log(f"\nunreadable images: {len(corrupt)}")
    for rel, err in corrupt[:15]:
        log(f"  {rel}  {err}")

    groups = {h: paths for h, paths in md5s.items() if len(paths) > 1}
    redundant = sum(len(v) - 1 for v in groups.values())
    log(f"\nexact-duplicate groups (MD5): {len(groups)}  ({redundant} redundant files)")
    for h, paths in sorted(groups.items(), key=lambda kv: -len(kv[1]))[:10]:
        log(f"  {h[:10]}  x{len(paths)}  e.g. {paths[0]}")

    log(f"\nmetadata files: {len(metas)}")
    for p in metas:
        log(f"  {p.relative_to(root)}")
        log(f"      {parse_check(p)}")
    if others:
        log(f"\nother files: {len(others)}")
        for p in others[:10]:
            log(f"  {p.relative_to(root)}")


def main():
    cfg = load_config()
    raw = PROJECT_ROOT / cfg["paths"]["raw"]
    out = PROJECT_ROOT / cfg["paths"]["processed"] / "new_datasets_inventory.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    names = sys.argv[1:] or ["mio", "piyarathne_oral"]

    lines = []

    def log(text=""):
        print(text, flush=True)
        lines.append(text)

    log("INVENTORY OF NEWLY ARRIVED DATASETS (data/raw is read-only)")
    for name in names:
        audit(raw / name, name, log)
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"\nsaved {out}")


if __name__ == "__main__":
    main()
